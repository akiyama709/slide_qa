#!/usr/bin/env python3
"""Optional independent-writer integration fixtures; installs nothing.
PPTX requires existing python-pptx; --format pdf requires existing ReportLab.
Creates only fictional files outside this repository.
"""
import argparse
from pathlib import Path

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output",type=Path)
    parser.add_argument("--format",choices=["pptx","pdf"],default="pptx")
    args=parser.parse_args()
    output=args.output.expanduser().resolve()
    root=Path(__file__).resolve().parents[1]
    if output==root or root in output.parents:
        raise SystemExit("Keep generated presentation fixtures outside the repository.")
    if output.exists():raise SystemExit("Refusing to overwrite an existing file.")
    output.parent.mkdir(parents=True,exist_ok=True)
    if args.format=="pdf":
        from reportlab.pdfgen.canvas import Canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
        pdf=Canvas(str(output),pagesize=(720,540))
        for title,lines in [
            ("Cobalt cooling",["The Cobalt chamber uses a water circuit."]),
            ("",[]),
            ("Museum lighting",["Museum lighting uses shielded lamps.","The design is provisional."]),
            ("温室の換気",["温室では気温が28度を超えたときに換気を行います．"])]:
            pdf.setFont("HeiseiKakuGo-W5" if title=="温室の換気" else "Helvetica",20)
            if title:pdf.drawString(40,460,title)
            for i,line in enumerate(lines):pdf.drawString(40,425-i*35,line)
            pdf.showPage()
        pdf.save()
        print("Created a fictional four-page PDF with English/Japanese text and an empty page.")
        return
    from pptx import Presentation
    presentation=Presentation()
    for title,body,notes in [
        ("Willow ventilation","Willow opens its vent at 28 C.","The vent opens only when humidity is below 70 percent."),
        ("Archive humidity","The archive humidity limit is 70 percent.","This limit is provisional and requires review.")]:
        slide=presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text=title;slide.placeholders[1].text=body
        slide.notes_slide.notes_text_frame.text=notes
    ids=presentation.slides._sldIdLst
    ids.insert(0,ids[1])
    presentation.save(output)
    print("Created a fictional two-slide PPTX; reordered relationship IDs and speaker notes included.")

if __name__=="__main__":main()
